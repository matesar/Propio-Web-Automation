#!/usr/bin/env python3
"""
Monitorea una tabla de historial de llamadas en una página ya abierta y logueada,
y guarda nuevas llamadas en Excel con cálculo de importe.

Backend principal: Playwright conectado a una sesión existente vía CDP.
"""

from __future__ import annotations

import argparse
import json
import sys
import time
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font
from openpyxl.styles import Alignment, Border, PatternFill, Side
from openpyxl.utils import get_column_letter
from playwright.sync_api import Browser, Error as PlaywrightError, Page, sync_playwright

# ---------------------------- Configuración global ---------------------------- #
RATE_PER_MINUTE = Decimal("0.11")
TABLE_ROW_SELECTOR = "mat-table.call-list-table mat-row"
EXCEL_HEADERS = [
    "Customer ID",
    "Call Date",
    "Call Start",
    "Duration (Minutes)",
    "Amount USD",
    "Detected At",
    "Unique Key",
]
SUMMARY_SHEET_NAME = "Daily Summary"
CALLS_SHEET_NAME = "Calls"
SUMMARY_HEADERS = [
    "Call Date",
    "Total Calls",
    "Total Minutes",
    "Total Amount USD",
    "Total Available Minutes",
]
STATUS_LOG_SHEET_NAME = "Status Log"
STATUS_DURATIONS_SHEET_NAME = "Status Durations"
STATUS_LOG_HEADERS = ["Interpreter ID", "Status", "Status Date", "Unique Status Key"]
STATUS_DURATION_HEADERS = [
    "Interpreter ID",
    "From Status",
    "To Status",
    "Start DateTime",
    "End DateTime",
    "Elapsed Seconds",
    "Elapsed Minutes",
    "Elapsed MM:SS",
    "Unique Duration Key",
]
MONTHLY_CALL_SUMMARY_SHEET_NAME = "Monthly Call Summary"
MONTHLY_CALL_SUMMARY_HEADERS = [
    "Year-Month",
    "Total Calls",
    "Total Minutes",
    "Total Amount USD",
]
WEEKLY_AVAILABLE_SUMMARY_SHEET_NAME = "Weekly Available Summary"
WEEKLY_AVAILABLE_SUMMARY_HEADERS = [
    "Week Start",
    "Week End",
    "Total Available Hours",
    "Target Hours",
    "Remaining Hours",
    "Over Target Hours",
    "Completion %",
]
WEEKLY_AVAILABLE_DETAIL_SHEET_NAME = "Weekly Available Detail"
WEEKLY_AVAILABLE_DETAIL_HEADERS = ["Week Start", "Date", "Day Name", "Available Hours"]

# Estilo visual (paleta sobria)
COLOR_HEADER_BG = "1F3A5F"   # azul oscuro
COLOR_HEADER_TXT = "FFFFFF"
COLOR_ZEBRA = "EEF3F8"       # gris azulado suave
COLOR_KPI_BG = "E8EEF6"
COLOR_GOOD = "D9EAD3"        # verde suave
COLOR_WARN = "FCE5CD"        # naranja suave
COLOR_BAD = "F4CCCC"         # rojo suave


@dataclass(frozen=True)
class CallRecord:
    """Representa una llamada normalizada lista para persistir."""

    customer_id: str
    call_date: str
    call_start: str
    duration_minutes: int

    @property
    def amount_usd(self) -> Decimal:
        return (Decimal(self.duration_minutes) * RATE_PER_MINUTE).quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP
        )

    @property
    def unique_key(self) -> str:
        return build_unique_key(
            self.customer_id, self.call_date, self.call_start, self.duration_minutes
        )


@dataclass(frozen=True)
class StatusEvent:
    interpreter_id: int
    status: str
    status_date: datetime

    @property
    def unique_key(self) -> str:
        return f"{self.interpreter_id}|{self.status}|{self.status_date.isoformat()}"


@dataclass(frozen=True)
class StatusDuration:
    interpreter_id: int
    from_status: str
    to_status: str
    start_dt: datetime
    end_dt: datetime

    @property
    def elapsed_seconds(self) -> int:
        return max(0, int((self.end_dt - self.start_dt).total_seconds()))

    @property
    def elapsed_minutes(self) -> Decimal:
        return (Decimal(self.elapsed_seconds) / Decimal("60")).quantize(Decimal("0.01"))

    @property
    def elapsed_mmss(self) -> str:
        mins, secs = divmod(self.elapsed_seconds, 60)
        return f"{mins:02d}:{secs:02d}"

    @property
    def unique_key(self) -> str:
        return (
            f"{self.interpreter_id}|{self.from_status}|{self.to_status}|"
            f"{self.start_dt.isoformat()}|{self.end_dt.isoformat()}"
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Monitorea nuevas llamadas en una tabla HTML y guarda en Excel."
    )
    parser.add_argument(
        "--cdp-url",
        default="http://127.0.0.1:9222",
        help="URL CDP del navegador abierto con remote debugging (default: http://127.0.0.1:9222)",
    )
    parser.add_argument(
        "--url-contains",
        default="",
        help="Texto para encontrar la pestaña correcta por URL (ej: /call-history). Si se omite, usa la pestaña activa encontrada.",
    )
    parser.add_argument(
        "--excel",
        default="call_log.xlsx",
        help="Ruta del archivo Excel de salida (default: call_log.xlsx)",
    )
    parser.add_argument(
        "--interval",
        type=int,
        default=180,
        help="Segundos entre cada polling (default: 180)",
    )
    return parser.parse_args()


# ------------------------- Conexión y selección de página ------------------------- #

def connect_browser_via_cdp(cdp_url: str) -> Tuple[object, Browser]:
    """Conecta Playwright a un navegador Chromium/Chrome/Edge ya abierto vía CDP."""
    pw = sync_playwright().start()
    try:
        browser = pw.chromium.connect_over_cdp(cdp_url)
    except Exception:
        pw.stop()
        raise
    return pw, browser


def find_target_page(browser: Browser, url_contains: str) -> Page:
    """Busca una página abierta que coincida con url_contains o toma la primera disponible."""
    candidates: List[Page] = []
    for context in browser.contexts:
        candidates.extend(context.pages)

    if not candidates:
        raise RuntimeError(
            "No se encontraron pestañas abiertas en la sesión conectada. "
            "Abrí la página de historial primero."
        )

    if url_contains:
        for p in candidates:
            if url_contains.lower() in p.url.lower():
                return p
        raise RuntimeError(
            f"No se encontró ninguna pestaña cuya URL contenga: {url_contains!r}."
        )

    # Si no se especifica filtro, usar la primera pestaña no vacía
    for p in candidates:
        if p.url and p.url != "about:blank":
            return p

    return candidates[0]


# ---------------------------- Lectura de tabla HTML ---------------------------- #

def wait_table_ready(page: Page, timeout_ms: int = 15_000) -> None:
    """Espera al menos una fila visible para confirmar que la tabla cargó."""
    page.wait_for_selector(TABLE_ROW_SELECTOR, timeout=timeout_ms)


def safe_text(locator) -> str:
    txt = locator.inner_text(timeout=2000)
    return " ".join(txt.split()).strip()


def parse_duration_to_int(raw: str) -> Optional[int]:
    """Convierte duración textual a entero de minutos (esperado: '5')."""
    cleaned = raw.strip()
    if not cleaned:
        return None
    # mantener solo dígitos por robustez ante espacios o sufijos accidentales
    digits = "".join(ch for ch in cleaned if ch.isdigit())
    if not digits:
        return None
    try:
        return int(digits)
    except ValueError:
        return None


def build_unique_key(customer_id: str, call_date: str, call_start: str, duration_minutes: int) -> str:
    return f"{customer_id}|{call_date}|{call_start}|{duration_minutes}"


def extract_record_from_row(row) -> Optional[CallRecord]:
    """
    Extrae datos de una fila mat-row.

    Estrategia de selectores:
    1) Preferir columnas por clase (más estable en Angular Material):
       - .cdk-column-customerId
       - .cdk-column-callDate
       - .cdk-column-callStart
       - .cdk-column-callDuration
    2) Fallback por índice de celdas si cambian clases pero no orden de columnas.
    """
    customer_id = ""
    call_date = ""
    call_start = ""
    duration_raw = ""

    # 1) Selectores de columna por clase
    col_map = {
        "customer_id": ".cdk-column-customerId, .mat-column-customerId",
        "call_date": ".cdk-column-callDate, .mat-column-callDate",
        "call_start": ".cdk-column-callStart, .mat-column-callStart",
        "duration": ".cdk-column-callDuration, .mat-column-callDuration",
    }

    if row.locator(col_map["customer_id"]).count() > 0:
        customer_id = safe_text(row.locator(col_map["customer_id"]).first)
    if row.locator(col_map["call_date"]).count() > 0:
        call_date = safe_text(row.locator(col_map["call_date"]).first)
    if row.locator(col_map["call_start"]).count() > 0:
        call_start = safe_text(row.locator(col_map["call_start"]).first)
    if row.locator(col_map["duration"]).count() > 0:
        duration_raw = safe_text(row.locator(col_map["duration"]).first)

    # 2) Fallback por índice (Customer ID, Call Date, Call Start, Duration=índice 3)
    if not (customer_id and call_date and call_start and duration_raw):
        cells = row.locator("mat-cell")
        if cells.count() >= 4:
            customer_id = customer_id or safe_text(cells.nth(0))
            call_date = call_date or safe_text(cells.nth(1))
            call_start = call_start or safe_text(cells.nth(2))
            duration_raw = duration_raw or safe_text(cells.nth(3))

    duration_minutes = parse_duration_to_int(duration_raw)
    if not customer_id or not call_date or not call_start or duration_minutes is None:
        return None

    return CallRecord(
        customer_id=customer_id,
        call_date=call_date,
        call_start=call_start,
        duration_minutes=duration_minutes,
    )


def read_visible_records(page: Page) -> List[CallRecord]:
    """Relee todas las filas visibles en cada ciclo (estrategia robusta)."""
    records: List[CallRecord] = []
    rows = page.locator(TABLE_ROW_SELECTOR)
    row_count = rows.count()

    for i in range(row_count):
        row = rows.nth(i)
        rec = extract_record_from_row(row)
        if rec:
            records.append(rec)
    return records


def read_visible_records_with_recovery(page: Page, retries: int = 1) -> List[CallRecord]:
    """
    Lee filas visibles con recuperación simple ante rerender/cambios transitorios del DOM.
    Reintenta relocalizando la tabla antes de propagar error.
    """
    last_error: Optional[Exception] = None
    for attempt in range(retries + 1):
        try:
            wait_table_ready(page, timeout_ms=10_000)
            return read_visible_records(page)
        except Exception as e:
            last_error = e
            if attempt < retries:
                print(f"[WARN] Error leyendo tabla. Reintentando relocalizar ({attempt + 1}/{retries})...")
                time.sleep(1)
    if last_error:
        raise last_error
    return []


# ------------------------------- Manejo de Excel ------------------------------- #

def apply_basic_sheet_formatting(
    ws,
    headers: List[str],
    amount_header: Optional[str] = None,
    numeric_headers: Optional[List[str]] = None,
    percentage_headers: Optional[List[str]] = None,
    date_headers: Optional[List[str]] = None,
    datetime_headers: Optional[List[str]] = None,
    center_headers: Optional[List[str]] = None,
    widths: Optional[Dict[str, int]] = None,
) -> None:
    """Aplica formato básico: header bold, freeze, autofilter y anchos."""
    bold = Font(bold=True, color=COLOR_HEADER_TXT)
    header_fill = PatternFill(fill_type="solid", fgColor=COLOR_HEADER_BG)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    for i, header in enumerate(headers, start=1):
        cell = ws.cell(1, i)
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        if widths and header in widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[header]

    ws.freeze_panes = "A2"
    max_col = len(headers)
    max_row = max(ws.max_row, 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

    # Zebra, bordes y alineación base
    for r in range(2, ws.max_row + 1):
        zebra_fill = PatternFill(fill_type="solid", fgColor=COLOR_ZEBRA) if r % 2 == 0 else None
        for c in range(1, max_col + 1):
            data_cell = ws.cell(r, c)
            data_cell.border = thin_border
            if zebra_fill:
                data_cell.fill = zebra_fill

    if amount_header:
        amount_col = get_column_index(ws, amount_header, -1)
        if amount_col > 0:
            for r in range(2, ws.max_row + 1):
                ws.cell(r, amount_col).number_format = "#,##0.00"
    if numeric_headers:
        for header_name in numeric_headers:
            col = get_column_index(ws, header_name, -1)
            if col > 0:
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, col).number_format = "#,##0.00"
                    ws.cell(r, col).alignment = Alignment(horizontal="right", vertical="center")
    if percentage_headers:
        for header_name in percentage_headers:
            col = get_column_index(ws, header_name, -1)
            if col > 0:
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, col).number_format = "0.00%"
                    ws.cell(r, col).alignment = Alignment(horizontal="right", vertical="center")
    if date_headers:
        for header_name in date_headers:
            col = get_column_index(ws, header_name, -1)
            if col > 0:
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, col).number_format = "yyyy-mm-dd"
                    ws.cell(r, col).alignment = Alignment(horizontal="center", vertical="center")
    if datetime_headers:
        for header_name in datetime_headers:
            col = get_column_index(ws, header_name, -1)
            if col > 0:
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, col).number_format = "yyyy-mm-dd hh:mm:ss"
                    ws.cell(r, col).alignment = Alignment(horizontal="center", vertical="center")
    if center_headers:
        for header_name in center_headers:
            col = get_column_index(ws, header_name, -1)
            if col > 0:
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, col).alignment = Alignment(horizontal="center", vertical="center")


def set_dashboard_title(ws, title: str, cell: str = "I2") -> None:
    ws[cell] = title
    ws[cell].font = Font(bold=True, size=18, color=COLOR_HEADER_BG)
    ws[cell].alignment = Alignment(horizontal="left", vertical="center")


def set_kpi_cards(ws, kpis: List[Tuple[str, str]], start_cell: str = "I4") -> None:
    start_col = ws[start_cell].column
    start_row = ws[start_cell].row
    card_fills = [COLOR_KPI_BG, COLOR_KPI_BG, COLOR_GOOD, COLOR_WARN]
    ws.column_dimensions[get_column_letter(start_col)].width = 28
    ws.column_dimensions[get_column_letter(start_col + 1)].width = 20
    for idx, (label, value) in enumerate(kpis):
        row = start_row + (idx * 2)
        label_cell = ws.cell(row=row, column=start_col, value=label)
        value_cell = ws.cell(row=row, column=start_col + 1, value=value)
        fill = PatternFill(fill_type="solid", fgColor=card_fills[idx % len(card_fills)])
        for c in (label_cell, value_cell):
            c.fill = fill
            c.border = Border(
                left=Side(style="thin", color="D9D9D9"),
                right=Side(style="thin", color="D9D9D9"),
                top=Side(style="thin", color="D9D9D9"),
                bottom=Side(style="thin", color="D9D9D9"),
            )
            c.alignment = Alignment(horizontal="center", vertical="center")
        label_cell.font = Font(bold=True, color=COLOR_HEADER_BG, size=12)
        value_cell.font = Font(bold=True, size=14)
        ws.row_dimensions[row].height = 26


def clear_sheet_charts(ws) -> None:
    ws._charts = []


def configure_workbook_layout(wb) -> None:
    """Configura visibilidad, orden de hojas y hoja activa."""
    hidden_sheet_names = [CALLS_SHEET_NAME, STATUS_LOG_SHEET_NAME, STATUS_DURATIONS_SHEET_NAME]
    for name in hidden_sheet_names:
        if name in wb.sheetnames:
            wb[name].sheet_state = "hidden"

    visible_order = [
        SUMMARY_SHEET_NAME,
        WEEKLY_AVAILABLE_SUMMARY_SHEET_NAME,
        WEEKLY_AVAILABLE_DETAIL_SHEET_NAME,
        MONTHLY_CALL_SUMMARY_SHEET_NAME,
    ]
    ordered = []
    used = set()
    for name in visible_order:
        if name in wb.sheetnames:
            ordered.append(wb[name])
            used.add(name)
    for ws in wb.worksheets:
        if ws.title not in used:
            ordered.append(ws)
    wb._sheets = ordered

    if SUMMARY_SHEET_NAME in wb.sheetnames:
        wb.active = wb.sheetnames.index(SUMMARY_SHEET_NAME)


def parse_datetime_value(value) -> Optional[datetime]:
    if isinstance(value, datetime):
        return value
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def parse_date_value(value) -> Optional[date]:
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def add_column_chart(
    ws,
    title: str,
    category_col: int,
    data_cols: List[int],
    anchor: str = "H8",
    y_axis_title: str = "",
) -> None:
    if ws.max_row < 2:
        return
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = title
    chart.y_axis.title = y_axis_title
    chart.x_axis.title = ""
    chart.height = 7
    chart.width = 12

    cats = Reference(ws, min_col=category_col, min_row=2, max_row=ws.max_row)
    for col in data_cols:
        data = Reference(ws, min_col=col, min_row=1, max_row=ws.max_row)
        chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, anchor)


def ensure_sheet_headers(ws, headers: List[str]) -> None:
    if ws.max_row < 1:
        ws.append(headers)
        return
    current = [ws.cell(1, i + 1).value for i in range(len(headers))]
    if current != headers:
        if ws.max_row == 1 and all(v is None for v in current):
            for i, h in enumerate(headers, start=1):
                ws.cell(1, i, h)
        else:
            print(f"[WARN] Encabezados en hoja '{ws.title}' no coinciden exactamente con lo esperado.")


def ensure_workbook(path: Path):
    """Crea/abre Excel y asegura todas las hojas de dataset + reporting."""
    if not path.exists():
        wb = Workbook()
        ws_calls = wb.active
        ws_calls.title = CALLS_SHEET_NAME
        ws_calls.append(EXCEL_HEADERS)
        ws_summary = wb.create_sheet(SUMMARY_SHEET_NAME)
        ws_summary.append(SUMMARY_HEADERS)
        ws_status_log = wb.create_sheet(STATUS_LOG_SHEET_NAME)
        ws_status_log.append(STATUS_LOG_HEADERS)
        ws_status_durations = wb.create_sheet(STATUS_DURATIONS_SHEET_NAME)
        ws_status_durations.append(STATUS_DURATION_HEADERS)
        ws_monthly_calls = wb.create_sheet(MONTHLY_CALL_SUMMARY_SHEET_NAME)
        ws_monthly_calls.append(MONTHLY_CALL_SUMMARY_HEADERS)
        ws_weekly_available = wb.create_sheet(WEEKLY_AVAILABLE_SUMMARY_SHEET_NAME)
        ws_weekly_available.append(WEEKLY_AVAILABLE_SUMMARY_HEADERS)
        ws_weekly_detail = wb.create_sheet(WEEKLY_AVAILABLE_DETAIL_SHEET_NAME)
        ws_weekly_detail.append(WEEKLY_AVAILABLE_DETAIL_HEADERS)
        apply_basic_sheet_formatting(
            ws_calls,
            EXCEL_HEADERS,
            amount_header="Amount USD",
            datetime_headers=["Detected At"],
            center_headers=["Call Date", "Call Start", "Duration (Minutes)"],
            widths={
                "Customer ID": 18,
                "Call Date": 14,
                "Call Start": 12,
                "Duration (Minutes)": 18,
                "Amount USD": 14,
                "Detected At": 20,
                "Unique Key": 50,
            },
        )
        apply_basic_sheet_formatting(
            ws_summary,
            SUMMARY_HEADERS,
            amount_header="Total Amount USD",
            numeric_headers=["Total Available Minutes"],
            date_headers=["Call Date"],
            center_headers=["Call Date", "Total Calls", "Total Minutes"],
            widths={
                "Call Date": 14,
                "Total Calls": 12,
                "Total Minutes": 14,
                "Total Amount USD": 18,
                "Total Available Minutes": 22,
            },
        )
        apply_basic_sheet_formatting(
            ws_status_log,
            STATUS_LOG_HEADERS,
            datetime_headers=["Status Date"],
            center_headers=["Interpreter ID", "Status"],
            widths={
                "Interpreter ID": 14,
                "Status": 10,
                "Status Date": 24,
                "Unique Status Key": 60,
            },
        )
        apply_basic_sheet_formatting(
            ws_status_durations,
            STATUS_DURATION_HEADERS,
            amount_header="Elapsed Minutes",
            datetime_headers=["Start DateTime", "End DateTime"],
            center_headers=["Interpreter ID", "From Status", "To Status", "Elapsed MM:SS"],
            widths={
                "Interpreter ID": 14,
                "From Status": 12,
                "To Status": 10,
                "Start DateTime": 24,
                "End DateTime": 24,
                "Elapsed Seconds": 16,
                "Elapsed Minutes": 16,
                "Elapsed MM:SS": 14,
                "Unique Duration Key": 70,
            },
        )
        apply_basic_sheet_formatting(
            ws_monthly_calls,
            MONTHLY_CALL_SUMMARY_HEADERS,
            amount_header="Total Amount USD",
            center_headers=["Year-Month", "Total Calls", "Total Minutes"],
            widths={
                "Year-Month": 14,
                "Total Calls": 12,
                "Total Minutes": 14,
                "Total Amount USD": 18,
            },
        )
        apply_basic_sheet_formatting(
            ws_weekly_available,
            WEEKLY_AVAILABLE_SUMMARY_HEADERS,
            numeric_headers=[
                "Total Available Hours",
                "Target Hours",
                "Remaining Hours",
                "Over Target Hours",
            ],
            percentage_headers=["Completion %"],
            date_headers=["Week Start", "Week End"],
            widths={
                "Week Start": 14,
                "Week End": 14,
                "Total Available Hours": 20,
                "Target Hours": 14,
                "Remaining Hours": 16,
                "Over Target Hours": 16,
                "Completion %": 14,
            },
        )
        apply_basic_sheet_formatting(
            ws_weekly_detail,
            WEEKLY_AVAILABLE_DETAIL_HEADERS,
            numeric_headers=["Available Hours"],
            date_headers=["Week Start", "Date"],
            center_headers=["Day Name"],
            widths={
                "Week Start": 14,
                "Date": 14,
                "Day Name": 12,
                "Available Hours": 16,
            },
        )
        configure_workbook_layout(wb)
        wb.save(path)
        return (
            wb,
            ws_calls,
            ws_summary,
            ws_status_log,
            ws_status_durations,
            ws_monthly_calls,
            ws_weekly_available,
            ws_weekly_detail,
        )

    wb = load_workbook(path)
    ws_calls = wb[CALLS_SHEET_NAME] if CALLS_SHEET_NAME in wb.sheetnames else wb.active
    if ws_calls.title != CALLS_SHEET_NAME:
        ws_calls.title = CALLS_SHEET_NAME
    ws_summary = (
        wb[SUMMARY_SHEET_NAME]
        if SUMMARY_SHEET_NAME in wb.sheetnames
        else wb.create_sheet(SUMMARY_SHEET_NAME)
    )
    ws_status_log = (
        wb[STATUS_LOG_SHEET_NAME]
        if STATUS_LOG_SHEET_NAME in wb.sheetnames
        else wb.create_sheet(STATUS_LOG_SHEET_NAME)
    )
    ws_status_durations = (
        wb[STATUS_DURATIONS_SHEET_NAME]
        if STATUS_DURATIONS_SHEET_NAME in wb.sheetnames
        else wb.create_sheet(STATUS_DURATIONS_SHEET_NAME)
    )
    ws_monthly_calls = (
        wb[MONTHLY_CALL_SUMMARY_SHEET_NAME]
        if MONTHLY_CALL_SUMMARY_SHEET_NAME in wb.sheetnames
        else wb.create_sheet(MONTHLY_CALL_SUMMARY_SHEET_NAME)
    )
    ws_weekly_available = (
        wb[WEEKLY_AVAILABLE_SUMMARY_SHEET_NAME]
        if WEEKLY_AVAILABLE_SUMMARY_SHEET_NAME in wb.sheetnames
        else wb.create_sheet(WEEKLY_AVAILABLE_SUMMARY_SHEET_NAME)
    )
    ws_weekly_detail = (
        wb[WEEKLY_AVAILABLE_DETAIL_SHEET_NAME]
        if WEEKLY_AVAILABLE_DETAIL_SHEET_NAME in wb.sheetnames
        else wb.create_sheet(WEEKLY_AVAILABLE_DETAIL_SHEET_NAME)
    )

    ensure_sheet_headers(ws_calls, EXCEL_HEADERS)
    ensure_sheet_headers(ws_summary, SUMMARY_HEADERS)
    ensure_sheet_headers(ws_status_log, STATUS_LOG_HEADERS)
    ensure_sheet_headers(ws_status_durations, STATUS_DURATION_HEADERS)
    ensure_sheet_headers(ws_monthly_calls, MONTHLY_CALL_SUMMARY_HEADERS)
    ensure_sheet_headers(ws_weekly_available, WEEKLY_AVAILABLE_SUMMARY_HEADERS)
    ensure_sheet_headers(ws_weekly_detail, WEEKLY_AVAILABLE_DETAIL_HEADERS)
    apply_basic_sheet_formatting(
        ws_calls,
        EXCEL_HEADERS,
        amount_header="Amount USD",
        datetime_headers=["Detected At"],
        center_headers=["Call Date", "Call Start", "Duration (Minutes)"],
        widths={
            "Customer ID": 18,
            "Call Date": 14,
            "Call Start": 12,
            "Duration (Minutes)": 18,
            "Amount USD": 14,
            "Detected At": 20,
            "Unique Key": 50,
        },
    )
    apply_basic_sheet_formatting(
        ws_summary,
        SUMMARY_HEADERS,
        amount_header="Total Amount USD",
        numeric_headers=["Total Available Minutes"],
        date_headers=["Call Date"],
        center_headers=["Call Date", "Total Calls", "Total Minutes"],
        widths={
            "Call Date": 14,
            "Total Calls": 12,
            "Total Minutes": 14,
            "Total Amount USD": 18,
            "Total Available Minutes": 22,
        },
    )
    apply_basic_sheet_formatting(
        ws_status_log,
        STATUS_LOG_HEADERS,
        datetime_headers=["Status Date"],
        center_headers=["Interpreter ID", "Status"],
        widths={
            "Interpreter ID": 14,
            "Status": 10,
            "Status Date": 24,
            "Unique Status Key": 60,
        },
    )
    apply_basic_sheet_formatting(
        ws_status_durations,
        STATUS_DURATION_HEADERS,
        amount_header="Elapsed Minutes",
        datetime_headers=["Start DateTime", "End DateTime"],
        center_headers=["Interpreter ID", "From Status", "To Status", "Elapsed MM:SS"],
        widths={
            "Interpreter ID": 14,
            "From Status": 12,
            "To Status": 10,
            "Start DateTime": 24,
            "End DateTime": 24,
            "Elapsed Seconds": 16,
            "Elapsed Minutes": 16,
            "Elapsed MM:SS": 14,
            "Unique Duration Key": 70,
        },
    )
    apply_basic_sheet_formatting(
        ws_monthly_calls,
        MONTHLY_CALL_SUMMARY_HEADERS,
        amount_header="Total Amount USD",
        center_headers=["Year-Month", "Total Calls", "Total Minutes"],
        widths={
            "Year-Month": 14,
            "Total Calls": 12,
            "Total Minutes": 14,
            "Total Amount USD": 18,
        },
    )
    apply_basic_sheet_formatting(
        ws_weekly_available,
        WEEKLY_AVAILABLE_SUMMARY_HEADERS,
        numeric_headers=[
            "Total Available Hours",
            "Target Hours",
            "Remaining Hours",
            "Over Target Hours",
        ],
        percentage_headers=["Completion %"],
        date_headers=["Week Start", "Week End"],
        widths={
            "Week Start": 14,
            "Week End": 14,
            "Total Available Hours": 20,
            "Target Hours": 14,
            "Remaining Hours": 16,
            "Over Target Hours": 16,
            "Completion %": 14,
        },
    )
    apply_basic_sheet_formatting(
        ws_weekly_detail,
        WEEKLY_AVAILABLE_DETAIL_HEADERS,
        numeric_headers=["Available Hours"],
        date_headers=["Week Start", "Date"],
        center_headers=["Day Name"],
        widths={
            "Week Start": 14,
            "Date": 14,
            "Day Name": 12,
            "Available Hours": 16,
        },
    )
    configure_workbook_layout(wb)
    wb.save(path)
    return (
        wb,
        ws_calls,
        ws_summary,
        ws_status_log,
        ws_status_durations,
        ws_monthly_calls,
        ws_weekly_available,
        ws_weekly_detail,
    )


def get_column_index(ws, header_name: str, fallback_index: int) -> int:
    """Obtiene índice 1-based de una columna por header; usa fallback si no la encuentra."""
    headers = [ws.cell(1, i + 1).value for i in range(ws.max_column)]
    for idx, header in enumerate(headers, start=1):
        if str(header).strip() == header_name:
            return idx
    return fallback_index


def load_existing_unique_keys(ws) -> Set[str]:
    unique_key_col = get_column_index(ws, "Unique Key", len(EXCEL_HEADERS))
    keys: Set[str] = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        key = row[unique_key_col - 1] if len(row) >= unique_key_col else None
        if key:
            keys.add(str(key))
    return keys


def calculate_total_amount(ws) -> Decimal:
    amount_col = get_column_index(
        ws, "Amount USD", EXCEL_HEADERS.index("Amount USD") + 1
    )
    total = Decimal("0")
    for row in ws.iter_rows(min_row=2, values_only=True):
        amount = row[amount_col - 1] if len(row) >= amount_col else None
        if amount is None:
            continue
        try:
            total += Decimal(str(amount))
        except (InvalidOperation, ValueError):
            continue
    return total.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def append_new_records(
    path: Path, ws_calls, wb, new_records_with_detection: List[Tuple[CallRecord, str]]
) -> Decimal:
    """Agrega registros nuevos al Excel y devuelve nuevo total acumulado."""
    for rec, detected_at in new_records_with_detection:
        detected_at_dt = parse_datetime_value(detected_at)
        ws_calls.append(
            [
                rec.customer_id,
                rec.call_date,
                rec.call_start,
                rec.duration_minutes,
                rec.amount_usd,
                detected_at_dt if detected_at_dt else detected_at,
                rec.unique_key,
            ]
        )
    apply_basic_sheet_formatting(
        ws_calls,
        EXCEL_HEADERS,
        amount_header="Amount USD",
        datetime_headers=["Detected At"],
        center_headers=["Call Date", "Call Start", "Duration (Minutes)"],
    )

    wb.save(path)
    return calculate_total_amount(ws_calls)


def aggregate_available_minutes_by_day(ws_status_durations) -> Dict[str, Decimal]:
    """
    Suma minutos Available por día desde Status Durations.
    Regla: sumar solo tramos con From Status = AV, agrupando por fecha de Start DateTime.
    """
    start_col = get_column_index(ws_status_durations, "Start DateTime", 4)
    end_col = get_column_index(ws_status_durations, "End DateTime", 5)
    from_status_col = get_column_index(ws_status_durations, "From Status", 2)
    grouped: Dict[str, Decimal] = {}

    for row in ws_status_durations.iter_rows(min_row=2, values_only=True):
        if len(row) < max(start_col, end_col, from_status_col):
            continue
        from_status = str(row[from_status_col - 1]).strip().upper()
        if from_status != "AV":
            continue
        try:
            start_dt = parse_datetime_value(row[start_col - 1])
            end_dt = parse_datetime_value(row[end_col - 1])
            if start_dt is None or end_dt is None:
                continue
        except (TypeError, ValueError):
            continue
        if end_dt < start_dt:
            continue
        day_key = start_dt.date().isoformat()
        minutes = (Decimal((end_dt - start_dt).total_seconds()) / Decimal("60")).quantize(
            Decimal("0.01")
        )
        grouped[day_key] = grouped.get(day_key, Decimal("0")) + minutes

    return grouped


def rebuild_daily_summary(ws_calls, ws_summary, ws_status_durations) -> Dict[str, Dict[str, Decimal]]:
    """Reconstruye Daily Summary combinando Calls y Status Durations."""
    call_date_col = get_column_index(ws_calls, "Call Date", 2)
    minutes_col = get_column_index(ws_calls, "Duration (Minutes)", 4)
    amount_col = get_column_index(ws_calls, "Amount USD", 5)

    daily: Dict[str, Dict[str, Decimal]] = {}
    for row in ws_calls.iter_rows(min_row=2, values_only=True):
        if len(row) < max(call_date_col, minutes_col, amount_col):
            continue
        parsed_call_date = parse_date_value(row[call_date_col - 1])
        if not parsed_call_date:
            continue
        call_date = parsed_call_date.isoformat()
        minutes_raw = row[minutes_col - 1]
        amount_raw = row[amount_col - 1]
        try:
            minutes = int(str(minutes_raw)) if minutes_raw is not None else 0
        except ValueError:
            minutes = 0
        try:
            amount = Decimal(str(amount_raw)) if amount_raw is not None else Decimal("0")
        except (InvalidOperation, ValueError):
            amount = Decimal("0")

        if call_date not in daily:
            daily[call_date] = {
                "total_calls": Decimal("0"),
                "total_minutes": Decimal("0"),
                "total_amount": Decimal("0"),
                "total_available_minutes": Decimal("0"),
            }
        daily[call_date]["total_calls"] += Decimal("1")
        daily[call_date]["total_minutes"] += Decimal(minutes)
        daily[call_date]["total_amount"] += amount

    available_by_day = aggregate_available_minutes_by_day(ws_status_durations)
    all_days = set(daily.keys()) | set(available_by_day.keys())
    for day in all_days:
        if day not in daily:
            daily[day] = {
                "total_calls": Decimal("0"),
                "total_minutes": Decimal("0"),
                "total_amount": Decimal("0"),
                "total_available_minutes": Decimal("0"),
            }
        daily[day]["total_available_minutes"] = available_by_day.get(day, Decimal("0"))

    ws_summary.delete_rows(2, ws_summary.max_row)
    for call_date in sorted(daily.keys()):
        call_date_out = parse_date_value(call_date) or call_date
        ws_summary.append(
            [
                call_date_out,
                int(daily[call_date]["total_calls"]),
                int(daily[call_date]["total_minutes"]),
                daily[call_date]["total_amount"].quantize(Decimal("0.01")),
                daily[call_date]["total_available_minutes"].quantize(Decimal("0.01")),
            ]
        )
    apply_basic_sheet_formatting(
        ws_summary,
        SUMMARY_HEADERS,
        amount_header="Total Amount USD",
        numeric_headers=["Total Available Minutes"],
        date_headers=["Call Date"],
        center_headers=["Call Date", "Total Calls", "Total Minutes"],
    )
    set_dashboard_title(ws_summary, "Daily Summary Dashboard")
    total_calls_all = sum(int(v["total_calls"]) for v in daily.values())
    total_minutes_all = sum(int(v["total_minutes"]) for v in daily.values())
    total_amount_all = sum(v["total_amount"] for v in daily.values()).quantize(Decimal("0.01"))
    total_available_hours_all = (
        sum(v["total_available_minutes"] for v in daily.values()) / Decimal("60")
    ).quantize(Decimal("0.01"))
    set_kpi_cards(
        ws_summary,
        [
            ("Total Calls", str(total_calls_all)),
            ("Total Minutes", str(total_minutes_all)),
            ("Total Amount USD", f"{total_amount_all:.2f}"),
            ("Total Available Hours", f"{total_available_hours_all:.2f}"),
        ],
        start_cell="H3",
    )
    clear_sheet_charts(ws_summary)
    return daily


def rebuild_monthly_call_summary(ws_calls, ws_monthly_calls) -> Dict[str, Dict[str, Decimal]]:
    """Reconstruye resumen mensual de llamadas desde Calls."""
    call_date_col = get_column_index(ws_calls, "Call Date", 2)
    minutes_col = get_column_index(ws_calls, "Duration (Minutes)", 4)
    amount_col = get_column_index(ws_calls, "Amount USD", 5)
    monthly: Dict[str, Dict[str, Decimal]] = {}

    for row in ws_calls.iter_rows(min_row=2, values_only=True):
        if len(row) < max(call_date_col, minutes_col, amount_col):
            continue
        parsed_call_date = parse_date_value(row[call_date_col - 1])
        if not parsed_call_date:
            continue
        month_start = parsed_call_date.replace(day=1)
        month_key = month_start.isoformat()

        try:
            minutes = int(str(row[minutes_col - 1])) if row[minutes_col - 1] is not None else 0
        except ValueError:
            minutes = 0
        try:
            amount = Decimal(str(row[amount_col - 1])) if row[amount_col - 1] is not None else Decimal("0")
        except (InvalidOperation, ValueError):
            amount = Decimal("0")

        if month_key not in monthly:
            monthly[month_key] = {
                "total_calls": Decimal("0"),
                "total_minutes": Decimal("0"),
                "total_amount": Decimal("0"),
            }
        monthly[month_key]["total_calls"] += Decimal("1")
        monthly[month_key]["total_minutes"] += Decimal(minutes)
        monthly[month_key]["total_amount"] += amount

    ws_monthly_calls.delete_rows(2, ws_monthly_calls.max_row)
    for month in sorted(monthly.keys()):
        month_date = datetime.fromisoformat(month).date()
        ws_monthly_calls.append(
            [
                month_date,
                int(monthly[month]["total_calls"]),
                int(monthly[month]["total_minutes"]),
                monthly[month]["total_amount"].quantize(Decimal("0.01")),
            ]
        )
    apply_basic_sheet_formatting(
        ws_monthly_calls,
        MONTHLY_CALL_SUMMARY_HEADERS,
        amount_header="Total Amount USD",
        center_headers=["Year-Month", "Total Calls", "Total Minutes"],
        date_headers=["Year-Month"],
    )
    month_col = get_column_index(ws_monthly_calls, "Year-Month", 1)
    for r in range(2, ws_monthly_calls.max_row + 1):
        ws_monthly_calls.cell(r, month_col).number_format = "mmm yyyy"
    set_dashboard_title(ws_monthly_calls, "Monthly Call Summary")
    clear_sheet_charts(ws_monthly_calls)
    add_column_chart(
        ws_monthly_calls,
        title="Total Amount USD por mes",
        category_col=1,
        data_cols=[4],
        anchor="F4",
        y_axis_title="USD",
    )
    return monthly


def get_week_start_sunday(dt: date) -> date:
    """Devuelve el domingo de la semana de dt (semana domingo-sábado)."""
    days_since_sunday = (dt.weekday() + 1) % 7
    return dt - timedelta(days=days_since_sunday)


def rebuild_weekly_available_summary(
    ws_status_durations, ws_weekly_available, ws_weekly_detail
) -> Dict[str, Dict[str, object]]:
    """
    Reconstruye resumen semanal (domingo-sábado) de Available en horas.
    Suma solo tramos con From Status = AV, usando Start DateTime para asignar semana.
    """
    start_col = get_column_index(ws_status_durations, "Start DateTime", 4)
    end_col = get_column_index(ws_status_durations, "End DateTime", 5)
    from_status_col = get_column_index(ws_status_durations, "From Status", 2)

    weekly: Dict[str, Dict[str, object]] = {}
    daily_detail: Dict[Tuple[str, str], Decimal] = {}
    for row in ws_status_durations.iter_rows(min_row=2, values_only=True):
        if len(row) < max(start_col, end_col, from_status_col):
            continue
        if str(row[from_status_col - 1]).strip().upper() != "AV":
            continue
        try:
            start_dt = parse_datetime_value(row[start_col - 1])
            end_dt = parse_datetime_value(row[end_col - 1])
            if start_dt is None or end_dt is None:
                continue
        except (TypeError, ValueError):
            continue
        if end_dt < start_dt:
            continue
        week_start = get_week_start_sunday(start_dt.date())
        week_end = week_start + timedelta(days=6)
        week_key = week_start.isoformat()
        hours = (Decimal((end_dt - start_dt).total_seconds()) / Decimal("3600")).quantize(
            Decimal("0.01")
        )

        if week_key not in weekly:
            weekly[week_key] = {
                "week_start_str": week_start.isoformat(),
                "week_end_str": week_end.isoformat(),
                "total_available_hours": Decimal("0"),
            }
        weekly[week_key]["total_available_hours"] = Decimal(
            str(weekly[week_key]["total_available_hours"])
        ) + hours
        day_key = start_dt.date().isoformat()
        daily_detail[(week_key, day_key)] = daily_detail.get((week_key, day_key), Decimal("0")) + hours

    ws_weekly_available.delete_rows(2, ws_weekly_available.max_row)
    target = Decimal("20.00")
    for week_key in sorted(weekly.keys()):
        total_hours = Decimal(str(weekly[week_key]["total_available_hours"])).quantize(
            Decimal("0.01")
        )
        remaining = max(Decimal("0"), target - total_hours).quantize(Decimal("0.01"))
        over_target = max(Decimal("0"), total_hours - target).quantize(Decimal("0.01"))
        ws_weekly_available.append(
            [
                datetime.fromisoformat(str(weekly[week_key]["week_start_str"])).date(),
                datetime.fromisoformat(str(weekly[week_key]["week_end_str"])).date(),
                total_hours,
                target,
                remaining,
                over_target,
                (total_hours / target).quantize(Decimal("0.0001")) if target > 0 else Decimal("0"),
            ]
        )
        weekly[week_key]["target_hours"] = target
        weekly[week_key]["remaining_hours"] = remaining
        weekly[week_key]["over_target_hours"] = over_target
        weekly[week_key]["completion_pct"] = (
            (total_hours / target).quantize(Decimal("0.0001")) if target > 0 else Decimal("0")
        )

    apply_basic_sheet_formatting(
        ws_weekly_available,
        WEEKLY_AVAILABLE_SUMMARY_HEADERS,
        numeric_headers=[
            "Total Available Hours",
            "Target Hours",
            "Remaining Hours",
            "Over Target Hours",
        ],
        percentage_headers=["Completion %"],
        date_headers=["Week Start", "Week End"],
    )
    # Formato condicional de objetivo semanal
    ws_weekly_available.conditional_formatting._cf_rules = {}
    rem_col = get_column_index(ws_weekly_available, "Remaining Hours", 5)
    over_col = get_column_index(ws_weekly_available, "Over Target Hours", 6)
    if ws_weekly_available.max_row >= 2:
        ws_weekly_available.conditional_formatting.add(
            f"{get_column_letter(rem_col)}2:{get_column_letter(rem_col)}{ws_weekly_available.max_row}",
            CellIsRule(
                operator="greaterThan",
                formula=["0"],
                fill=PatternFill(fill_type="solid", fgColor=COLOR_WARN),
            ),
        )
        ws_weekly_available.conditional_formatting.add(
            f"{get_column_letter(rem_col)}2:{get_column_letter(rem_col)}{ws_weekly_available.max_row}",
            CellIsRule(
                operator="greaterThan",
                formula=["5"],
                fill=PatternFill(fill_type="solid", fgColor=COLOR_BAD),
            ),
        )
        ws_weekly_available.conditional_formatting.add(
            f"{get_column_letter(over_col)}2:{get_column_letter(over_col)}{ws_weekly_available.max_row}",
            CellIsRule(
                operator="greaterThan",
                formula=["0"],
                fill=PatternFill(fill_type="solid", fgColor=COLOR_GOOD),
            ),
        )
    set_dashboard_title(ws_weekly_available, "Weekly Available Summary")
    clear_sheet_charts(ws_weekly_available)
    add_column_chart(
        ws_weekly_available,
        title="Available Hours por semana",
        category_col=1,
        data_cols=[3, 4],
        anchor="H8",
        y_axis_title="Hours",
    )

    # Detalle diario (domingo a sábado) para gráfico secundario
    ws_weekly_detail.delete_rows(2, ws_weekly_detail.max_row)
    day_names = ["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]
    for week_key in sorted({k[0] for k in daily_detail.keys()}):
        week_start = datetime.fromisoformat(week_key).date()
        for i in range(7):
            day_date = week_start + timedelta(days=i)
            hours = daily_detail.get((week_key, day_date.isoformat()), Decimal("0")).quantize(
                Decimal("0.01")
            )
            ws_weekly_detail.append([week_start, day_date, day_names[i], hours])
    apply_basic_sheet_formatting(
        ws_weekly_detail,
        WEEKLY_AVAILABLE_DETAIL_HEADERS,
        numeric_headers=["Available Hours"],
        date_headers=["Week Start", "Date"],
        center_headers=["Day Name"],
    )

    # gráfico diario: toma la última semana cargada
    clear_sheet_charts(ws_weekly_detail)
    if ws_weekly_detail.max_row >= 8:
        last_week_start = ws_weekly_detail.cell(ws_weekly_detail.max_row, 1).value
        if isinstance(last_week_start, date):
            start_row = None
            end_row = None
            for r in range(2, ws_weekly_detail.max_row + 1):
                if ws_weekly_detail.cell(r, 1).value == last_week_start and start_row is None:
                    start_row = r
                if start_row is not None and ws_weekly_detail.cell(r, 1).value != last_week_start:
                    end_row = r - 1
                    break
            if start_row is not None and end_row is None:
                end_row = ws_weekly_detail.max_row
            if start_row is not None and end_row is not None and end_row >= start_row:
                chart = BarChart()
                chart.type = "col"
                chart.style = 10
                chart.title = f"Daily Available Hours ({last_week_start.isoformat()} week)"
                chart.y_axis.title = "Hours"
                chart.height = 7
                chart.width = 12
                data = Reference(ws_weekly_detail, min_col=4, min_row=start_row - 1, max_row=end_row)
                cats = Reference(ws_weekly_detail, min_col=3, min_row=start_row, max_row=end_row)
                chart.add_data(data, titles_from_data=True)
                chart.set_categories(cats)
                ws_weekly_detail.add_chart(chart, "F8")
    return weekly


def get_today_summary(
    daily_summary: Dict[str, Dict[str, Decimal]]
) -> Tuple[int, int, Decimal, Decimal]:
    """Obtiene resumen del día local actual basado en claves de Call Date parseables."""
    today = date.today()
    total_calls = 0
    total_minutes = 0
    total_amount = Decimal("0")
    total_available_minutes = Decimal("0")

    for call_date, values in daily_summary.items():
        parsed_date: Optional[date] = None
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
            try:
                parsed_date = datetime.strptime(call_date[:10], fmt).date()
                break
            except ValueError:
                continue
        if parsed_date == today:
            total_calls += int(values["total_calls"])
            total_minutes += int(values["total_minutes"])
            total_amount += values["total_amount"]
            total_available_minutes += values.get("total_available_minutes", Decimal("0"))

    return (
        total_calls,
        total_minutes,
        total_amount.quantize(Decimal("0.01")),
        total_available_minutes.quantize(Decimal("0.01")),
    )


def get_current_month_and_week_metrics(
    monthly_summary: Dict[str, Dict[str, Decimal]],
    weekly_available_summary: Dict[str, Dict[str, object]],
) -> Tuple[int, Decimal, Decimal, Decimal]:
    """
    Devuelve:
    - total_calls_mes_actual
    - total_amount_mes_actual
    - total_available_hours_semana_actual
    - remaining_hours_semana_actual
    """
    today = date.today()
    month_key = today.replace(day=1).isoformat()
    week_start = get_week_start_sunday(today).isoformat()

    month_calls = (
        int(monthly_summary[month_key]["total_calls"]) if month_key in monthly_summary else 0
    )
    month_amount = (
        monthly_summary[month_key]["total_amount"].quantize(Decimal("0.01"))
        if month_key in monthly_summary
        else Decimal("0.00")
    )

    if week_start in weekly_available_summary:
        week_row = weekly_available_summary[week_start]
        available_hours = Decimal(str(week_row["total_available_hours"])).quantize(
            Decimal("0.01")
        )
        remaining = Decimal(str(week_row.get("remaining_hours", Decimal("20.00")))).quantize(
            Decimal("0.01")
        )
    else:
        available_hours = Decimal("0.00")
        remaining = Decimal("20.00")

    return month_calls, month_amount, available_hours, remaining


def parse_status_payload(payload: Dict) -> Optional[StatusEvent]:
    """Parsea payload de status y devuelve StatusEvent si cumple estructura."""
    required = {"interpreterId", "status", "statusDate"}
    if not isinstance(payload, dict) or not required.issubset(payload.keys()):
        return None
    try:
        interpreter_id = int(payload["interpreterId"])
        status = str(payload["status"]).strip().upper()
        status_date = datetime.fromisoformat(str(payload["statusDate"]))
    except (ValueError, TypeError):
        return None
    if not status:
        return None
    return StatusEvent(interpreter_id=interpreter_id, status=status, status_date=status_date)


def register_status_request_listener(page: Page, pending_status_events: List[StatusEvent]) -> None:
    """Escucha requests de red y acumula eventos de estado parseables."""

    def on_request(request) -> None:
        payload = None
        try:
            payload = request.post_data_json
        except Exception:
            raw = request.post_data
            if raw:
                try:
                    payload = json.loads(raw)
                except json.JSONDecodeError:
                    payload = None
        event = parse_status_payload(payload) if payload else None
        if event:
            pending_status_events.append(event)

    page.on("request", on_request)


def load_status_keys_and_last_states(ws_status_log) -> Tuple[Set[str], Dict[int, StatusEvent]]:
    seen_status_keys: Set[str] = set()
    last_state_by_interpreter: Dict[int, StatusEvent] = {}
    i_id_col = get_column_index(ws_status_log, "Interpreter ID", 1)
    status_col = get_column_index(ws_status_log, "Status", 2)
    status_date_col = get_column_index(ws_status_log, "Status Date", 3)
    key_col = get_column_index(ws_status_log, "Unique Status Key", 4)

    for row in ws_status_log.iter_rows(min_row=2, values_only=True):
        if len(row) < max(i_id_col, status_col, status_date_col, key_col):
            continue
        key = row[key_col - 1]
        if key:
            seen_status_keys.add(str(key))
        try:
            interpreter_id = int(row[i_id_col - 1])
            status = str(row[status_col - 1]).strip().upper()
            status_dt = parse_datetime_value(row[status_date_col - 1])
            if status_dt is None:
                continue
        except (TypeError, ValueError):
            continue
        candidate = StatusEvent(interpreter_id=interpreter_id, status=status, status_date=status_dt)
        prev = last_state_by_interpreter.get(interpreter_id)
        if prev is None or candidate.status_date >= prev.status_date:
            last_state_by_interpreter[interpreter_id] = candidate
    return seen_status_keys, last_state_by_interpreter


def load_status_duration_keys(ws_status_durations) -> Set[str]:
    seen_duration_keys: Set[str] = set()
    key_col = get_column_index(ws_status_durations, "Unique Duration Key", 9)
    for row in ws_status_durations.iter_rows(min_row=2, values_only=True):
        if len(row) >= key_col and row[key_col - 1]:
            seen_duration_keys.add(str(row[key_col - 1]))
    return seen_duration_keys


def build_status_duration(prev_event: StatusEvent, new_event: StatusEvent) -> Optional[StatusDuration]:
    if new_event.status_date < prev_event.status_date:
        return None
    return StatusDuration(
        interpreter_id=new_event.interpreter_id,
        from_status=prev_event.status,
        to_status=new_event.status,
        start_dt=prev_event.status_date,
        end_dt=new_event.status_date,
    )


def process_pending_status_events(
    pending_status_events: List[StatusEvent],
    ws_status_log,
    ws_status_durations,
    seen_status_keys: Set[str],
    seen_duration_keys: Set[str],
    last_state_by_interpreter: Dict[int, StatusEvent],
) -> int:
    """
    Procesa eventos pendientes:
    - ignora duplicados exactos por unique status key
    - guarda solo cambios reales de estado
    - calcula y guarda duración de transición entre estado previo y nuevo
    """
    if not pending_status_events:
        return 0

    processed = 0
    while pending_status_events:
        event = pending_status_events.pop(0)
        if event.unique_key in seen_status_keys:
            continue

        prev = last_state_by_interpreter.get(event.interpreter_id)
        if prev and prev.status == event.status:
            seen_status_keys.add(event.unique_key)
            continue

        ws_status_log.append(
            [
                event.interpreter_id,
                event.status,
                event.status_date,
                event.unique_key,
            ]
        )
        seen_status_keys.add(event.unique_key)
        processed += 1

        if prev:
            duration = build_status_duration(prev, event)
            if duration and duration.unique_key not in seen_duration_keys:
                ws_status_durations.append(
                    [
                        duration.interpreter_id,
                        duration.from_status,
                        duration.to_status,
                        duration.start_dt,
                        duration.end_dt,
                        duration.elapsed_seconds,
                        duration.elapsed_minutes,
                        duration.elapsed_mmss,
                        duration.unique_key,
                    ]
                )
                seen_duration_keys.add(duration.unique_key)
                if duration.from_status == "AV" and duration.to_status == "LO":
                    print(
                        "[STATUS AV->LO] "
                        f"interpreter={duration.interpreter_id} "
                        f"inicio={duration.start_dt.isoformat()} "
                        f"fin={duration.end_dt.isoformat()} "
                        f"segundos={duration.elapsed_seconds} "
                        f"minutos={duration.elapsed_minutes} "
                        f"mm:ss={duration.elapsed_mmss}"
                    )

        last_state_by_interpreter[event.interpreter_id] = event

    apply_basic_sheet_formatting(
        ws_status_log,
        STATUS_LOG_HEADERS,
        datetime_headers=["Status Date"],
        center_headers=["Interpreter ID", "Status"],
    )
    apply_basic_sheet_formatting(
        ws_status_durations,
        STATUS_DURATION_HEADERS,
        amount_header="Elapsed Minutes",
        datetime_headers=["Start DateTime", "End DateTime"],
        center_headers=["Interpreter ID", "From Status", "To Status", "Elapsed MM:SS"],
    )
    return processed


# ------------------------------- Lógica principal ------------------------------- #

def print_new_call(rec: CallRecord, detected_at: str, total_amount: Decimal) -> None:
    print("\n[NEW CALL] Nueva llamada detectada")
    print(f"  Customer ID        : {rec.customer_id}")
    print(f"  Call Date          : {rec.call_date}")
    print(f"  Call Start         : {rec.call_start}")
    print(f"  Duration (Minutes) : {rec.duration_minutes}")
    print(f"  Detected At        : {detected_at}")
    print(f"  Amount USD         : {rec.amount_usd}")
    print(f"  Total acumulado    : {total_amount}")


def print_initial_sync_call(rec: CallRecord, detected_at: str, total_amount: Decimal) -> None:
    print("\n[INITIAL SYNC] Llamada visible importada")
    print(f"  Customer ID        : {rec.customer_id}")
    print(f"  Call Date          : {rec.call_date}")
    print(f"  Call Start         : {rec.call_start}")
    print(f"  Duration (Minutes) : {rec.duration_minutes}")
    print(f"  Detected At        : {detected_at}")
    print(f"  Amount USD         : {rec.amount_usd}")
    print(f"  Total acumulado    : {total_amount}")


def monitor_loop(page: Page, excel_path: Path, interval_seconds: int) -> None:
    (
        wb,
        ws_calls,
        ws_summary,
        ws_status_log,
        ws_status_durations,
        ws_monthly_calls,
        ws_weekly_available,
        ws_weekly_detail,
    ) = ensure_workbook(excel_path)
    seen_keys = load_existing_unique_keys(ws_calls)
    persisted_count_before_sync = len(seen_keys)
    seen_status_keys, last_state_by_interpreter = load_status_keys_and_last_states(ws_status_log)
    seen_duration_keys = load_status_duration_keys(ws_status_durations)
    pending_status_events: List[StatusEvent] = []
    register_status_request_listener(page, pending_status_events)

    # Sincronización inicial: importar llamadas visibles no presentes en Excel
    visible_records = read_visible_records_with_recovery(page, retries=1)
    initial_sync_records: List[CallRecord] = [
        rec for rec in visible_records if rec.unique_key not in seen_keys
    ]
    initial_sync_with_detection = [
        (rec, time.strftime("%Y-%m-%d %H:%M:%S")) for rec in initial_sync_records
    ]
    initial_imported_count = len(initial_sync_with_detection)

    if initial_sync_with_detection:
        total_amount = append_new_records(
            excel_path, ws_calls, wb, initial_sync_with_detection
        )
        for rec, _ in initial_sync_with_detection:
            seen_keys.add(rec.unique_key)
    else:
        total_amount = calculate_total_amount(ws_calls)

    daily_summary = rebuild_daily_summary(ws_calls, ws_summary, ws_status_durations)
    monthly_summary = rebuild_monthly_call_summary(ws_calls, ws_monthly_calls)
    weekly_available_summary = rebuild_weekly_available_summary(
        ws_status_durations, ws_weekly_available, ws_weekly_detail
    )
    configure_workbook_layout(wb)
    wb.save(excel_path)
    today_calls, today_minutes, today_amount, today_available_minutes = get_today_summary(
        daily_summary
    )
    month_calls, month_amount, week_available_hours, week_remaining_hours = (
        get_current_month_and_week_metrics(monthly_summary, weekly_available_summary)
    )

    print("=" * 72)
    print("Monitor iniciado")
    print(f"Archivo Excel         : {excel_path.resolve()}")
    print(f"Registros ya guardados en Excel (antes de sync): {persisted_count_before_sync}")
    print(f"Sincronización inicial: {initial_imported_count} llamadas visibles importadas")
    print(f"Total acumulado general USD (post-sync): {total_amount}")
    print(
        f"Resumen hoy (post-sync) -> llamadas: {today_calls} | minutos: {today_minutes} | "
        f"importe USD: {today_amount} | available_minutes_hoy: {today_available_minutes}"
    )
    print(
        f"Resumen mes/semana actual -> calls_mes: {month_calls} | "
        f"importe_mes_usd: {month_amount} | available_hours_semana: {week_available_hours} | "
        f"remaining_hours_semana(20h): {week_remaining_hours}"
    )
    print(f"Polling cada          : {interval_seconds} segundos")
    print("Presioná Ctrl+C para detener.")
    print("=" * 72)
    print(
        f"Status monitor inicializado | eventos status ya guardados: {len(seen_status_keys)} | "
        f"tramos guardados: {len(seen_duration_keys)}"
    )
    for rec, detected_at in initial_sync_with_detection:
        print_initial_sync_call(rec, detected_at, total_amount)
    if initial_imported_count:
        print(f"[INITIAL SYNC] Completado: {initial_imported_count} llamadas visibles importadas.")
    else:
        print("[INITIAL SYNC] Completado: no hubo llamadas visibles nuevas para importar.")

    while True:
        cycle_start = time.strftime("%Y-%m-%d %H:%M:%S")
        try:
            status_updates = process_pending_status_events(
                pending_status_events=pending_status_events,
                ws_status_log=ws_status_log,
                ws_status_durations=ws_status_durations,
                seen_status_keys=seen_status_keys,
                seen_duration_keys=seen_duration_keys,
                last_state_by_interpreter=last_state_by_interpreter,
            )
            if status_updates:
                daily_summary = rebuild_daily_summary(ws_calls, ws_summary, ws_status_durations)
                monthly_summary = rebuild_monthly_call_summary(ws_calls, ws_monthly_calls)
                weekly_available_summary = rebuild_weekly_available_summary(
                    ws_status_durations, ws_weekly_available, ws_weekly_detail
                )
                today_calls, today_minutes, today_amount, today_available_minutes = get_today_summary(
                    daily_summary
                )
                month_calls, month_amount, week_available_hours, week_remaining_hours = (
                    get_current_month_and_week_metrics(
                        monthly_summary, weekly_available_summary
                    )
                )
                configure_workbook_layout(wb)
                wb.save(excel_path)
                print(f"[STATUS] Cambios reales de estado guardados: {status_updates}")
                print(
                    f"  Resumen hoy -> llamadas: {today_calls} | minutos: {today_minutes} | "
                    f"importe USD: {today_amount} | available_minutes_hoy: {today_available_minutes}"
                )
                print(
                    f"  Resumen mes/semana -> calls_mes: {month_calls} | "
                    f"importe_mes_usd: {month_amount} | available_hours_semana: {week_available_hours} | "
                    f"remaining_hours_semana(20h): {week_remaining_hours}"
                )

            all_records = read_visible_records_with_recovery(page, retries=1)
            new_records: List[CallRecord] = []

            # Detectar nuevos por unique key contra lo ya persistido
            for rec in all_records:
                if rec.unique_key not in seen_keys:
                    new_records.append(rec)

            if new_records:
                new_records_with_detection = [
                    (rec, time.strftime("%Y-%m-%d %H:%M:%S")) for rec in new_records
                ]
                total_amount = append_new_records(
                    excel_path, ws_calls, wb, new_records_with_detection
                )
                daily_summary = rebuild_daily_summary(ws_calls, ws_summary, ws_status_durations)
                monthly_summary = rebuild_monthly_call_summary(ws_calls, ws_monthly_calls)
                weekly_available_summary = rebuild_weekly_available_summary(
                    ws_status_durations, ws_weekly_available, ws_weekly_detail
                )
                configure_workbook_layout(wb)
                wb.save(excel_path)
                today_calls, today_minutes, today_amount, today_available_minutes = get_today_summary(
                    daily_summary
                )
                month_calls, month_amount, week_available_hours, week_remaining_hours = (
                    get_current_month_and_week_metrics(
                        monthly_summary, weekly_available_summary
                    )
                )
                for rec, detected_at in new_records_with_detection:
                    seen_keys.add(rec.unique_key)
                    print_new_call(rec, detected_at, total_amount)
                print(
                    f"  Resumen hoy -> llamadas: {today_calls} | minutos: {today_minutes} | "
                    f"importe USD: {today_amount} | available_minutes_hoy: {today_available_minutes}"
                )
                print(
                    f"  Resumen mes/semana -> calls_mes: {month_calls} | "
                    f"importe_mes_usd: {month_amount} | available_hours_semana: {week_available_hours} | "
                    f"remaining_hours_semana(20h): {week_remaining_hours}"
                )
            else:
                print(
                    f"[{cycle_start}] Sin novedades. Filas visibles leídas: {len(all_records)} | "
                    f"Total acumulado USD: {total_amount}"
                )

        except PlaywrightError as e:
            print(f"[ERROR] Fallo Playwright durante lectura: {e}")
        except Exception as e:
            print(f"[ERROR] Excepción inesperada: {e}")

        time.sleep(interval_seconds)


def main() -> int:
    args = parse_args()
    excel_path = Path(args.excel)

    pw = None
    browser = None
    try:
        pw, browser = connect_browser_via_cdp(args.cdp_url)
        page = find_target_page(browser, args.url_contains)

        print(f"Conectado a pestaña: {page.url}")
        wait_table_ready(page)

        monitor_loop(page, excel_path, args.interval)
        return 0

    except KeyboardInterrupt:
        print("\nDetenido por usuario.")
        return 0
    except Exception as e:
        print(f"[FATAL] {e}")
        return 1
    finally:
        if pw:
            pw.stop()


if __name__ == "__main__":
    sys.exit(main())
